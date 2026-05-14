import sys
import os
import csv
from openpyxl import load_workbook
import datetime

def limpar_valor(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    v_str = str(v).strip().replace('R$', '').replace('\xa0', '').strip()
    if not v_str:
        return 0.0
    # Caso string brasileira "3.150,50" -> 3150.50
    # Nota: se não tiver vírgula e tiver ponto "3150.50", tentamos manter seguro.
    if ',' in v_str:
        v_str = v_str.replace('.', '').replace(',', '.')
    try:
        return float(v_str)
    except:
        return 0.0

def format_date(d):
    if isinstance(d, datetime.datetime):
        return d.strftime("%d/%m/%Y")
    # Caso venha como string "13/05/2026"
    return str(d).strip()

def is_valid_date(d):
    return isinstance(d, datetime.datetime) or (isinstance(d, str) and len(d) >= 8 and '/' in d)

def sanitize_nasajon(input_path, output_path):
    print(f"Lendo o extrato Nasajon: {input_path} ... Pode demorar alguns segundos.")
    
    wb = load_workbook(filename=input_path, data_only=True)
    ws = wb.active
    
    linhas_mescladas = []
    linha_atual = None
    
    linhas_originais = list(ws.iter_rows(values_only=True))
    
    # 1. Passo de varredura para identificar e juntar históricos de múltiplas linhas 
    for idx, row in enumerate(linhas_originais):
        if not row or all(v is None for v in row):
            continue
            
        col_A = row[0] # Data (ou cabeçalho / vazio em linhas de continuação)
        
        if is_valid_date(col_A):
            # Encontrou nova linha principal, se havia uma pendente, salvamos na lista de consolidadas
            if linha_atual is not None:
                linhas_mescladas.append(linha_atual)
                
            # Cria a base da linha. Nasajon tem no mínimo 11 colunas no padrao (K é a 11ª = index 10)
            linha_base = list(row)
            # Garantir 11 colunas na lista da linha pra evitar index out of bounds
            while len(linha_base) < 11:
                linha_base.append(None)
                
            linha_atual = linha_base
        else:
            # Não é linha de cabeçalho nem data. Pode ser linha que "sobra" dados que não cabiam em cima.
            # O Nasajon gosta de quebrar o 'Histórico' na coluna C (index 2)
            if linha_atual is not None and len(row) > 2 and row[2]:
                historico_extra = str(row[2]).strip()
                if historico_extra and historico_extra != 'None':
                    # Concatena no historico da linha_atual
                    atual_hist = str(linha_atual[2] or '').strip()
                    if atual_hist and atual_hist != 'None':
                        linha_atual[2] = f"{atual_hist} {historico_extra}"
                    else:
                        linha_atual[2] = historico_extra

    # Gravar a última lida
    if linha_atual is not None:
        linhas_mescladas.append(linha_atual)
        
    print(f"Total de registros mesclados contabilizados: {len(linhas_mescladas)}")
    
    # 2. Passo de filtragem e exportação para CSV Mastigado
    registros_validos = []
    
    for row in linhas_mescladas:
        # Puxa conforme especificado pelas colunas
        data = format_date(row[0])
        documento = str(row[1] or '').strip()
        historico = str(row[2] or '').strip()
        info_conciliado = str(row[3] or '').strip() # Somente Lançamentos Conciliados
        classif_financeira = str(row[4] or '').replace(' ', '')
        centro_custo = str(row[5] or '').strip()
        # row[6] é pulada (Branco)
        despesa = limpar_valor(row[7])
        receita = limpar_valor(row[8])
        saldo = limpar_valor(row[9])
        sinal_saldo = str(row[10] or '').strip().upper() # C ou D
        
        # Ignorar desnecessarios/linhas que não sejam entradas válidas (Header perdido)
        if not classif_financeira and not historico:
            continue
            
        # D: Informação de Conciliado. Se não possuir NENHUM conteúdo que simbolize conciliamento.. 
        # (Não sei exato qual string, mas se for vazio ele não exporta. Se for False/nao, ignora)
        # Por garantia, só importaremos se a coluna não estiver em 'vazia' ou com conteúdo negativo ("Não", "Cancelado", etc)
        if not info_conciliado or info_conciliado.lower() in ['não', 'nao', 'n', 'false', 'cancelado', 'none']:
            continue
            
        # Deduzimos o tipo central ("DESPESA" ou "RECEITA")
        # Se for preenchido H, é Despesa. Se I, Receita.
        if despesa > 0:
            tipo = "DESPESA"
            valor = despesa
        elif receita > 0:
            tipo = "RECEITA"
            valor = receita
        else:
            # Pode ocorrer linhas de banco malucas com zeros? Exporta como Outro
            tipo = "OUTROS"
            valor = 0.0
            
        # Registramos a row mestigada e limpinha
        registros_validos.append({
            'Data': data,
            'Documento': documento,
            'Historico': historico,
            'Classificacao': classif_financeira,
            'CentroCusto': centro_custo,
            'Tipo': tipo,
            'Valor': f"{valor:.2f}",
            'SaldoFinal': f"{saldo:.2f}",
            'SinalSaldo': sinal_saldo
        })
        
    print(f"Lançamentos conciliados limpos e preparados: {len(registros_validos)}")
    
    # 3. Exportando CSV Padronizado (delimitador Ponto e Vírgula)
    with open(output_path, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=[
            'Data', 'Documento', 'Historico', 'Classificacao', 'CentroCusto', 'Tipo', 'Valor', 'SaldoFinal', 'SinalSaldo'
        ], delimiter=';')
        
        writer.writeheader()
        writer.writerows(registros_validos)
        
    print(f"CSV criado com sucesso e higienizado: {output_path} ! Está pronto!")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        input_test_msg = "\n📌 Como usar no terminal:\npython sanitizador_nasajon.py <sua_planilha_nasajon.xlsx> <arquivo_saida.csv>\n"
        print(input_test_msg)
        sys.exit(1)
        
    arquivo_entrada = sys.argv[1]
    arquivo_saida = sys.argv[2]
    
    if not os.path.exists(arquivo_entrada):
        print(f"Arquivo não encontrado: {arquivo_entrada}")
        sys.exit(1)
        
    sanitize_nasajon(arquivo_entrada, arquivo_saida)
